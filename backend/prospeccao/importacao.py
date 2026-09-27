"""Importação em lote de revendedoras a partir de planilha Excel (PRD fase 2).

Reaproveita as mesmas regras do wizard: `criar_revendedora` exige telefone,
ao menos um interesse e marca quando "vende outras marcas". A validação
acontece linha a linha **antes** de gravar, então a prévia mostra todos os
erros da planilha sem tocar no banco.
"""
from __future__ import annotations

import datetime
import io
import re
import unicodedata
from dataclasses import dataclass, field

from django.core.exceptions import ValidationError
from django.core.validators import validate_email
from localflavor.br.br_states import STATE_CHOICES
from openpyxl import Workbook, load_workbook

from .models import ProdutoInteresse, Revendedora
from .services import criar_revendedora, obter_ou_criar_cidade

MAX_LINHAS = 1000
SEPARADOR = re.compile(r"[;,]")

COLUNAS = [
    "nome_completo",
    "telefone",
    "cpf",
    "cidade",
    "uf",
    "interesses",
    "vende_outras_marcas",
    "marcas",
    "email",
    "data_nascimento",
    "cep",
    "endereco",
    "bairro",
    "rg",
]

OBRIGATORIAS = ["nome_completo", "telefone", "cidade", "uf", "interesses"]

EXEMPLO = {
    "nome_completo": "Maria da Silva",
    "telefone": "81999990001",
    "cpf": "12345678901",
    "cidade": "Recife",
    "uf": "PE",
    "interesses": "Cosméticos; Vitaminas",
    "vende_outras_marcas": "sim",
    "marcas": "Natura; Avon",
    "email": "maria@exemplo.com",
    "data_nascimento": "15/03/1990",
    "cep": "50010000",
    "endereco": "Rua das Flores, 120",
    "bairro": "Boa Vista",
    "rg": "1234567",
}

ALIAS = {
    "nome": "nome_completo",
    "nome_completo": "nome_completo",
    "revendedora": "nome_completo",
    "telefone": "telefone",
    "whatsapp": "telefone",
    "celular": "telefone",
    "fone": "telefone",
    "telefone_whatsapp": "telefone",
    "cpf": "cpf",
    "rg": "rg",
    "email": "email",
    "e_mail": "email",
    "data_nascimento": "data_nascimento",
    "nascimento": "data_nascimento",
    "dt_nascimento": "data_nascimento",
    "cep": "cep",
    "endereco": "endereco",
    "bairro": "bairro",
    "cidade": "cidade",
    "municipio": "cidade",
    "uf": "uf",
    "estado": "uf",
    "vende_outras_marcas": "vende_outras_marcas",
    "vende_outras": "vende_outras_marcas",
    "outras_marcas": "vende_outras_marcas",
    "marcas": "marcas",
    "marca": "marcas",
    "interesses": "interesses",
    "interesse": "interesses",
    "produtos": "interesses",
    "categorias": "interesses",
}

SIM = {"sim", "s", "1", "x", "true", "verdadeiro"}
NAO = {"nao", "n", "0", "f", "false", "falso"}
UFS = {uf for uf, _ in STATE_CHOICES}


class ErroDePlanilha(Exception):
    """Planilha ilegível (sem colunas obrigatórias, vazia ou grande demais)."""


def normaliza_cabecalho(valor) -> str:
    texto = unicodedata.normalize("NFKD", str(valor or "")).encode("ascii", "ignore").decode()
    texto = re.sub(r"[^0-9a-zA-Z]+", "_", texto.strip().lower())
    return texto.strip("_")


def _texto(valor) -> str:
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    if isinstance(valor, datetime.datetime):
        valor = valor.date()
    if isinstance(valor, datetime.date):
        return valor.strftime("%d/%m/%Y")
    return str(valor).strip()


def _digitos(valor: str) -> str:
    return "".join(ch for ch in valor if ch.isdigit())


def _lista(valor: str) -> list[str]:
    return [item.strip() for item in SEPARADOR.split(valor) if item.strip()]


def ler_planilha(arquivo) -> list[dict[str, str]]:
    """Lê a primeira aba e devolve as linhas como `{coluna: texto}`."""
    try:
        pasta = load_workbook(io.BytesIO(arquivo.read()), read_only=True, data_only=True)
    except Exception as erro:  # noqa: BLE001 - openpyxl lança tipos variados
        raise ErroDePlanilha(f"Não consegui abrir o arquivo: {erro}") from erro

    aba = pasta.worksheets[0]
    linhas = aba.iter_rows(values_only=True)
    try:
        cabecalho = next(linhas)
    except StopIteration as erro:
        raise ErroDePlanilha("A planilha está vazia.") from erro

    mapeamento = {}
    for posicao, bruto in enumerate(cabecalho):
        chave = ALIAS.get(normaliza_cabecalho(bruto))
        if chave and chave not in mapeamento:
            mapeamento[chave] = posicao

    ausentes = [coluna for coluna in OBRIGATORIAS if coluna not in mapeamento]
    if ausentes:
        raise ErroDePlanilha(
            "Colunas obrigatórias ausentes: " + ", ".join(ausentes) + "."
            " Baixe o modelo para conferir o formato."
        )

    saida = []
    for numero, célula in enumerate(linhas, start=2):
        if all(valor is None for valor in célula):
            continue
        registro = {
            coluna: _texto(célula[posicao])
            for coluna, posicao in mapeamento.items()
            if posicao < len(célula)
        }
        if any(registro.values()):
            registro["_linha"] = str(numero)
            saida.append(registro)

    if not saida:
        raise ErroDePlanilha("A planilha não tem nenhuma linha de dados.")
    if len(saida) > MAX_LINHAS:
        raise ErroDePlanilha(f"Máximo de {MAX_LINHAS} linhas por importação.")
    return saida


@dataclass
class ResultadoLinha:
    """Uma linha da planilha já validada e pronta para gravar."""

    numero: int
    original: dict[str, str]
    erros: list[str] = field(default_factory=list)
    dados: dict = field(default_factory=dict)
    telefones: list[dict] = field(default_factory=list)
    marcas: list[str] | None = None
    interesses: list[int] = field(default_factory=list)
    resumo: dict = field(default_factory=dict)

    @property
    def ok(self) -> bool:
        return not self.erros


def _valida_email(valor: str) -> str | None:
    if not valor:
        return None
    try:
        validate_email(valor)
    except ValidationError:
        return "e-mail inválido"
    return None


def _valida_data(valor: str) -> tuple[datetime.date | None, str | None]:
    if not valor:
        return None, None
    for formato in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
        try:
            data = datetime.datetime.strptime(valor, formato).date()
        except ValueError:
            continue
        if data > datetime.date.today():
            return None, "data de nascimento não pode ser futura"
        return data, None
    return None, "data inválida (use DD/MM/AAAA)"


def _valida_vende_outras(valor: str) -> tuple[bool | None, str | None]:
    limpo = unicodedata.normalize("NFKD", valor).encode("ascii", "ignore").decode().lower()
    if not limpo:
        return False, None
    if limpo in SIM:
        return True, None
    if limpo in NAO:
        return False, None
    return None, 'use "sim" ou "não"'


def _resolve_interesses(valor: str, catalogo: dict[str, int], erros: list[str]) -> list[int]:
    nomes = _lista(valor)
    if not nomes:
        erros.append("informe ao menos um produto de interesse")
        return []
    ids = []
    for nome in nomes:
        chave = normaliza_cabecalho(nome)
        if chave not in catalogo:
            erros.append(f"interesse desconhecido: {nome}")
        else:
            ids.append(catalogo[chave])
    return ids


def _catalogo_interesses() -> tuple[dict[str, int], dict[int, str]]:
    """`(nome_norm -> id, id -> nome)` dos produtos ativos, buscados uma vez só."""
    ativos = list(ProdutoInteresse.objects.filter(ativa=True).order_by("ordem", "nome"))
    return (
        {normaliza_cabecalho(p.nome): p.pk for p in ativos},
        {p.pk: p.nome for p in ativos},
    )


def analisa(linhas: list[dict[str, str]]) -> list[ResultadoLinha]:
    """Valida cada linha e monta os dados no formato do `criar_revendedora`."""
    catalogo, nomes_por_id = _catalogo_interesses()
    cpfs_no_banco = set(Revendedora.objects.filter(cpf__isnull=False).values_list("cpf", flat=True))

    resultados: list[ResultadoLinha] = []
    cpfs_na_planilha: dict[str, int] = {}

    for linha in linhas:
        numero = int(linha.get("_linha") or 0)
        original = {k: v for k, v in linha.items() if k != "_linha"}
        erros: list[str] = []

        nome = original.get("nome_completo", "")
        if not nome:
            erros.append("nome completo obrigatório")

        telefone = _digitos(original.get("telefone", ""))
        if not telefone:
            erros.append("telefone obrigatório")
        elif not 10 <= len(telefone) <= 11:
            erros.append("telefone precisa ter DDD + número (10 ou 11 dígitos)")

        cidade = original.get("cidade", "")
        uf = original.get("uf", "").upper()
        if not cidade:
            erros.append("cidade obrigatória")
        if uf not in UFS:
            erros.append(f"UF inválida: {uf or '—'}")

        cpf = _digitos(original.get("cpf", ""))
        if cpf:
            if len(cpf) != 11:
                erros.append("CPF precisa ter 11 dígitos")
            elif cpf in cpfs_no_banco:
                erros.append("CPF já cadastrado no sistema")
            elif cpf in cpfs_na_planilha:
                erros.append(f"CPF repetido na planilha (linha {cpfs_na_planilha[cpf]})")
            else:
                cpfs_na_planilha[cpf] = numero

        data_nascimento, erro_data = _valida_data(original.get("data_nascimento", ""))
        if erro_data:
            erros.append(erro_data)

        cep = _digitos(original.get("cep", ""))
        if original.get("cep") and len(cep) != 8:
            erros.append("CEP precisa ter 8 dígitos")

        erro_email = _valida_email(original.get("email", ""))
        if erro_email:
            erros.append(erro_email)

        vende_outras, erro_vende = _valida_vende_outras(original.get("vende_outras_marcas", ""))
        if erro_vende:
            erros.append(erro_vende)

        marcas = _lista(original.get("marcas", ""))
        if vende_outras and not marcas:
            erros.append('marque "sim" e informe as marcas, ou marque "não"')
        if not vende_outras:
            marcas = []

        interesses = _resolve_interesses(original.get("interesses", ""), catalogo, erros)

        resultado = ResultadoLinha(
            numero=numero,
            original=original,
            erros=erros,
            dados={
                "nome_completo": nome,
                "cpf": cpf or None,
                "rg": original.get("rg", ""),
                "email": original.get("email", ""),
                "data_nascimento": data_nascimento,
                "cep": cep,
                "endereco": original.get("endereco", ""),
                "bairro": original.get("bairro", ""),
                "vende_outras_marcas": bool(vende_outras),
                "cidade_nome": cidade,
                "cidade_uf": uf,
            },
            telefones=[{"numero": telefone, "tipo": "whatsapp", "principal": True}],
            marcas=marcas if vende_outras else None,
            interesses=interesses,
            resumo={
                "nome": nome or "—",
                "telefone": telefone or "—",
                "cidade": f"{cidade}/{uf}" if cidade and uf in UFS else "—",
                "interesses": ", ".join(nomes_por_id.get(pk, "") for pk in interesses),
                "marcas": ", ".join(marcas) if marcas else "—",
                "ja_revende": "sim" if marcas else "não",
            },
        )
        resultados.append(resultado)

    return resultados


def importa(resultados: list[ResultadoLinha], promotora) -> dict:
    """Grava as linhas válidas. Retorna `{criadas, puladas, falhas}`."""
    criadas = 0
    falhas: list[str] = []
    for resultado in resultados:
        if not resultado.ok:
            continue
        dados = dict(resultado.dados)
        dados["cidade"] = obter_ou_criar_cidade(dados.pop("cidade_nome"), dados.pop("cidade_uf"))
        try:
            criar_revendedora(
                promotora=promotora,
                dados=dados,
                telefones=resultado.telefones,
                marcas=resultado.marcas,
                interesses=resultado.interesses,
            )
            criadas += 1
        except ValidationError as erro:
            falhas.append(f"Linha {resultado.numero}: {'; '.join(erro.messages)}")

    return {
        "criadas": criadas,
        "puladas": sum(1 for r in resultados if not r.ok),
        "falhas": falhas,
    }


def modelo() -> bytes:
    """Planilha-modelo com as colunas esperadas e duas linhas de exemplo."""
    pasta = Workbook()
    aba = pasta.active
    aba.title = "revendedoras"
    aba.append(COLUNAS)
    aba.append([EXEMPLO[c] for c in COLUNAS])
    segunda = dict(EXEMPLO)
    segunda.update(
        {"nome_completo": "Joana Oliveira", "telefone": "81999990002", "cpf": "12345678902",
         "cidade": "Olinda", "uf": "PE", "vende_outras_marcas": "não", "marcas": "",
         "interesses": "Alimentos", "data_nascimento": "02/07/1985"}
    )
    aba.append([segunda[c] for c in COLUNAS])
    for coluna, largura in {"A": 24, "B": 16, "C": 14, "E": 6, "F": 26, "G": 20, "H": 22}.items():
        aba.column_dimensions[coluna].width = largura

    buffer = io.BytesIO()
    pasta.save(buffer)
    return buffer.getvalue()
